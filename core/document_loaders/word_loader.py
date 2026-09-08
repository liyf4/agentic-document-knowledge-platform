from __future__ import annotations

from pathlib import Path
from typing import Iterator, List, Tuple

from langchain_core.documents import Document


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    return "\n".join([
        "| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |",
        *("| " + " | ".join(row) + " |" for row in rows[1:]),
    ])


def _iter_body_items(docx) -> Iterator[Tuple[str, object]]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in docx.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield "paragraph", Paragraph(child, docx)
        elif child.tag.endswith("}tbl"):
            yield "table", Table(child, docx)


def _heading_level(style: str) -> int:
    lowered = style.casefold()
    if lowered.startswith("heading"):
        suffix = "".join(ch for ch in lowered if ch.isdigit())
        return max(1, int(suffix or 1))
    if lowered.startswith("标题"):
        suffix = "".join(ch for ch in lowered if ch.isdigit())
        return max(1, int(suffix or 1))
    return 0


def _core_properties(docx) -> dict:
    props = docx.core_properties
    result = {}
    for key in ("title", "author", "subject", "creator", "created", "modified"):
        value = getattr(props, key, None)
        if value:
            target = {"created": "creation_date", "modified": "modification_date"}.get(key, key)
            result[f"document_{target}"] = str(value)
    return result


def load_docx_documents(file_path: str, file_name: str) -> List[Document]:
    from docx import Document as DocxDocument

    docx = DocxDocument(file_path)
    docs: List[Document] = []
    headings: List[str] = []
    properties = _core_properties(docx)
    body_index = 0
    for kind, item in _iter_body_items(docx):
        if kind == "paragraph":
            text = item.text.strip()
            if not text:
                continue
            style = item.style.name if item.style else ""
            level = _heading_level(style)
            is_toc = style.casefold().startswith("toc") or style.startswith("目录")
            if level:
                headings = headings[:level - 1] + [text]
                chunk_kind = "heading"
            else:
                chunk_kind = "toc" if is_toc else "paragraph"
            tier = "secondary" if is_toc else "primary"
            docs.append(Document(page_content=text, metadata={
                "source": file_name, "file_type": "docx", "chunk_kind": chunk_kind,
                "page_role": "toc" if is_toc else "body", "retrieval_tier": tier,
                "paragraph_index": body_index, "style": style,
                "section_path": " > ".join(headings), "heading_level": level,
                "pre_chunked": True, **properties,
            }))
        else:
            table_text = _table_to_markdown(item)
            if table_text:
                docs.append(Document(page_content=table_text, metadata={
                    "source": file_name, "file_type": "docx", "chunk_kind": "table",
                    "page_role": "body", "retrieval_tier": "primary", "table_index": body_index,
                    "section_path": " > ".join(headings), "pre_chunked": True, **properties,
                }))
        body_index += 1

    seen = set()
    for section_index, section in enumerate(docx.sections):
        for kind, container in (("header", section.header), ("footer", section.footer)):
            text = "\n".join(paragraph.text.strip() for paragraph in container.paragraphs if paragraph.text.strip())
            key = (kind, text)
            if not text or key in seen:
                continue
            seen.add(key)
            docs.append(Document(page_content=text, metadata={
                "source": file_name, "file_type": "docx", "chunk_kind": kind,
                "page_role": "body", "retrieval_tier": "secondary", "section_index": section_index,
                "section_path": "", "pre_chunked": True, **properties,
            }))
    if not docs:
        raise ValueError(f"{Path(file_name).name} contains no readable text.")
    return docs
